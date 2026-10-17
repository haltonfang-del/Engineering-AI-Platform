from pathlib import Path
from app.knowledge.schemas import DocumentChunk


def parse_pptx(path: Path, relative_path: str) -> list[DocumentChunk]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    chunks: list[DocumentChunk] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    texts.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    row_text = " | ".join(value for value in cells if value)
                    if row_text:
                        texts.append(row_text)
        content = "\n".join(texts).strip()
        if content:
            chunks.append(DocumentChunk(
                source_name=path.name, source_path=relative_path, file_type=path.suffix.lower(),
                content=content, slide_number=slide_number,
            ))
    return chunks
