from pathlib import Path
from openpyxl import load_workbook
from app.knowledge.schemas import DocumentChunk


def parse_xlsx(path: Path, relative_path: str) -> list[DocumentChunk]:
    formula_book = load_workbook(path, read_only=True, data_only=False)
    value_book = load_workbook(path, read_only=True, data_only=True)
    chunks: list[DocumentChunk] = []
    try:
        for sheet_name in formula_book.sheetnames:
            formula_sheet = formula_book[sheet_name]
            value_sheet = value_book[sheet_name]
            for row in formula_sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if value is None:
                        continue
                    display = str(value)
                    if isinstance(value, str) and value.startswith("="):
                        cached = value_sheet[cell.coordinate].value
                        display = f"Formula: {value}"
                        if cached is not None:
                            display += f" | Cached value: {cached}"
                    elif not str(value).strip():
                        continue
                    chunks.append(DocumentChunk(
                        source_name=path.name, source_path=relative_path,
                        file_type=path.suffix.lower(),
                        content=f"Sheet: {sheet_name}\nCell: {cell.coordinate}\nValue: {display}",
                        sheet_name=sheet_name, cell_reference=cell.coordinate,
                    ))
    finally:
        formula_book.close()
        value_book.close()
    return chunks
