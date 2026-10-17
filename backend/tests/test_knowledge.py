from pathlib import Path

import pytest
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from reportlab.pdfgen import canvas

from app.knowledge.parsers.dispatcher import parse_document
from app.knowledge.service import KnowledgeService


def make_pptx(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text = "Spacer torque design"
    presentation.save(path)


def make_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "DesignData"
    sheet["A1"] = "Torque"
    sheet["B1"] = 12.5
    sheet["C1"] = "=B1*2"
    workbook.save(path)


def make_text_pdf(path: Path) -> None:
    document = canvas.Canvas(str(path))
    document.drawString(72, 720, "Bearing load specification")
    document.showPage()
    document.drawString(72, 720, "Second page torque limits")
    document.save()


def make_blank_pdf(path: Path) -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with path.open("wb") as stream:
        writer.write(stream)


def test_pptx_parser_keeps_slide_number(tmp_path: Path) -> None:
    path = tmp_path / "sample.pptx"
    make_pptx(path)
    chunks = parse_document(path, "sample.pptx")
    assert chunks
    assert chunks[0].slide_number == 1
    assert "Spacer torque design" in chunks[0].content


def test_xlsx_parser_keeps_sheet_and_cell(tmp_path: Path) -> None:
    path = tmp_path / "sample.xlsx"
    make_xlsx(path)
    chunks = parse_document(path, "sample.xlsx")
    torque = next(chunk for chunk in chunks if chunk.cell_reference == "A1")
    formula = next(chunk for chunk in chunks if chunk.cell_reference == "C1")
    assert torque.sheet_name == "DesignData"
    assert "Torque" in torque.content
    assert "Formula: =B1*2" in formula.content


def test_pdf_parser_keeps_page_number(tmp_path: Path) -> None:
    path = tmp_path / "text.pdf"
    make_text_pdf(path)
    chunks = parse_document(path, "text.pdf")
    assert len(chunks) == 2
    assert chunks[0].page_number == 1
    assert chunks[1].page_number == 2
    assert "Bearing load specification" in chunks[0].content


def test_blank_pdf_returns_no_text_chunks(tmp_path: Path) -> None:
    path = tmp_path / "blank.pdf"
    make_blank_pdf(path)
    assert parse_document(path, "blank.pdf") == []


def test_keyword_search_returns_source_metadata(tmp_path: Path) -> None:
    knowledge = tmp_path / "knowledge"
    knowledge.mkdir()
    make_pptx(knowledge / "sample.pptx")
    service = KnowledgeService(knowledge, tmp_path / "index.json")
    assert service.index_file("sample.pptx") == 1
    results = service.search("SPACER")
    assert results
    assert results[0].slide_number == 1
    assert results[0].source_path == "sample.pptx"


def test_path_traversal_is_rejected(tmp_path: Path) -> None:
    knowledge = tmp_path / "knowledge"
    knowledge.mkdir()
    make_xlsx(tmp_path / "outside.xlsx")
    service = KnowledgeService(knowledge, tmp_path / "index.json")
    with pytest.raises(ValueError):
        service.index_file("../outside.xlsx")


def test_unsupported_file_type_is_rejected(tmp_path: Path) -> None:
    knowledge = tmp_path / "knowledge"
    knowledge.mkdir()
    (knowledge / "notes.txt").write_text("not supported", encoding="utf-8")
    service = KnowledgeService(knowledge, tmp_path / "index.json")
    with pytest.raises(ValueError, match="Unsupported"):
        service.index_file("notes.txt")
